from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(
    r"C:\Users\PC\Documents\GitHub\nl-am-requirement-parser\实习报告_自然语言驱动智能化增材制造系统.docx"
)

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(96, 104, 116)
BLACK = RGBColor(0, 0, 0)


def set_run_font(
    run,
    *,
    ascii_name: str = "Calibri",
    east_asia: str = "宋体",
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: RGBColor | None = None,
) -> None:
    run.font.name = ascii_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_name)
    rfonts.set(qn("w:hAnsi"), ascii_name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_style_font(style, *, ascii_name: str, east_asia: str, size: float) -> None:
    style.font.name = ascii_name
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_name)
    rfonts.set(qn("w:hAnsi"), ascii_name)
    rfonts.set(qn("w:eastAsia"), east_asia)


def add_page_number(paragraph) -> None:
    paragraph.add_run("实习报告  |  第 ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instr, separate, value, end])
    paragraph.add_run(" 页")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    set_style_font(normal, ascii_name="Calibri", east_asia="宋体", size=11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.widow_control = True

    heading_specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[name]
        set_style_font(style, ascii_name="Calibri", east_asia="微软雅黑", size=size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    if "Abstract" not in doc.styles:
        abstract_style = doc.styles.add_style("Abstract", WD_STYLE_TYPE.PARAGRAPH)
    else:
        abstract_style = doc.styles["Abstract"]
    set_style_font(abstract_style, ascii_name="Calibri", east_asia="宋体", size=10.5)
    abstract_style.font.color.rgb = RGBColor(40, 48, 58)
    abstract_style.paragraph_format.space_after = Pt(6)
    abstract_style.paragraph_format.line_spacing = 1.20
    abstract_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract_style.paragraph_format.first_line_indent = Pt(21)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("自然语言驱动的智能化增材制造系统开发")
    set_run_font(hr, east_asia="微软雅黑", size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    add_page_number(fp)
    for run in fp.runs:
        set_run_font(run, east_asia="宋体", size=8.5, color=MUTED)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    r = kicker.add_run("实  习  报  告")
    set_run_font(r, east_asia="微软雅黑", size=14, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    r = title.add_run("自然语言驱动的智能化\n增材制造系统开发")
    set_run_font(r, east_asia="微软雅黑", size=24, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(54)
    r = subtitle.add_run("——“物质创制”AI 融合 3D 打印项目")
    set_run_font(r, east_asia="微软雅黑", size=13, color=DARK_BLUE)

    metadata = [
        "姓      名：____________________",
        "学      号：____________________",
        "专业班级：____________________",
        "实习单位：____________________",
        "指导教师：____________________",
        "实习时间：____________________",
    ]
    for line in metadata:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(line)
        set_run_font(r, east_asia="宋体", size=12, color=BLACK)

    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True


def add_body(doc: Document, text: str, *, first_indent: bool = True) -> None:
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Pt(22) if first_indent else Pt(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.widow_control = True
    r = p.add_run(text)
    set_run_font(r, east_asia="宋体", size=11, color=BLACK)


def add_abstract(doc: Document) -> None:
    add_heading(doc, "摘  要", 1)
    abstract = (
        "本次实习围绕“物质创制——自然语言驱动的智能化增材制造”项目展开。我参与搭建了从中文自然语言或本地语音输入，"
        "到需求结构化、三维模型生成与修复、尺寸归一化、Bambu Studio 自动定向与切片，再到文件交付、打印机上传和运行状态监测的完整流程。"
        "项目的重点不是简单串联若干工具，而是通过模块契约、状态持久化、文件哈希和安全边界，使 AI 生成结果能够进入可验证、可恢复的制造过程。"
        "实习中我完成了需求解析、模型处理、自动化工作流、本地网页交互、语音转写、设备连接以及测试文档等工作，并针对自然语言歧义、网格缺陷、外部工具调用、"
        "多材料映射和实体打印不可逆等问题进行了持续调试。"
    )
    p = doc.add_paragraph(style="Abstract")
    r = p.add_run(abstract)
    set_run_font(r, east_asia="宋体", size=10.5, color=RGBColor(40, 48, 58))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    label = p.add_run("关键词：")
    set_run_font(label, east_asia="微软雅黑", size=10.5, bold=True, color=DARK_BLUE)
    r = p.add_run("自然语言处理；生成式 AI；3D 打印；自动切片；Bambu X1C；工程自动化")
    set_run_font(r, east_asia="宋体", size=10.5, color=BLACK)


SECTIONS: list[tuple[str, int, list[str]]] = [
    (
        "一、实习背景与项目目标",
        1,
        [
            "传统 3D 打印通常需要用户先学习建模软件，再手动导入切片器、设置材料和支撑参数，最后通过设备端完成打印。对非专业用户而言，这条流程环节多、门槛高，任何一步出错都可能导致模型无法打印。项目提出“语音输入—AI 生成模型—自动切片—直接打印—检测优化”的新范式，希望让用户用一句自然语言描述需求，系统就能将想法逐步转化为可制造文件，并在明确授权后连接真实打印机。",
            "我的实习目标是把这一产品设想落实为可运行的软件原型。具体来说，一方面要完成自然语言理解、AI 建模和制造执行之间的数据贯通；另一方面要保证每个阶段都有清晰输入输出、失败原因和恢复方式。由于系统最终可能触发实体设备，我还把安全性和可追溯性作为核心要求：默认只生成文件，只有用户显式开启自动打印时才允许上传和启动，无法确认结果的打印动作不得自动重放。",
        ],
    ),
    (
        "二、主要工作内容",
        1,
        [],
    ),
    (
        "（一）需求解析与模块化架构",
        2,
        [
            "我首先完善了 M1 需求理解模块，将中文描述转换为经过校验的结构化 JSON。系统能够区分创意摆件和工程零件，提取对象、功能、姿态、尺寸、材料及制造约束，并统一毫米、厘米、米以及牛、千牛等单位。对于缺少尺寸、载荷或接口信息的工程需求，程序不会让大模型自行补全，而是生成澄清问题并阻止任务进入下一阶段。为此我设计了 JSON Schema、状态字段和 Manifest 交付契约，使后续模块只读取已确认的规格，而不再重复解释原始文本。",
            "在整体架构上，我将系统划分为需求解析、模型生成、制造准备、设备执行和前端交互等模块，并用统一工作流调度。每个阶段开始、完成或失败时都会写入状态文件和追加式事件日志，已经完成的确定性步骤可以复用。这样的设计使长时间运行的任务在程序中断后仍能恢复，也方便定位故障发生在语义解析、模型提供器、网格处理、切片还是打印机连接。",
        ],
    ),
    (
        "（二）三维模型生成、网格检查与尺寸恢复",
        2,
        [
            "在 M2 阶段，我接入了本地 TripoSG、Meshy 及离线 Mock 等模型提供方式，并实现了 Windows 与 WSL 环境之间的任务桥接。模型返回后，我没有直接把文件交给切片器，而是先检查 GLB/STL 的文件身份、三角面、连通分量、边界、法向、体积和包围盒，再根据问题选择清理坏面、保留主体、表面修复或体素修复。修复过程设有预算和验收条件，防止为了“修好”网格而过度改变外形。",
            "随后我完成了尺寸归一化和 STL 交付回读验证。系统根据用户确认的目标高度计算缩放比例，统一坐标系与单位，并在导出后重新读取文件，核对尺寸、哈希和记录是否一致。通过这项工作，我认识到生成式 AI 给出的三维结果只能算候选资产，只有经过几何检查、修复、尺度恢复和证据记录后，才能成为后续制造环节可信的输入。",
        ],
    ),
    (
        "（三）自动切片、材料配置与打印执行",
        2,
        [
            "在制造准备阶段，我调用 Bambu Studio 的后台能力实现 Auto Orient，并强制使用原生 tree(auto) 或 tree_hybrid 支撑策略，自动输出可编辑工程文件、G-code 3MF 和本体 STL。系统把切片器是否成功生成目标文件作为制造准备结果，同时保留几何文件和切片文件的哈希，网页下载时再次核验路径与内容，避免跨任务误取文件。对于旧任务，我还设计了只复用已确认 M1/M2 结果、重新执行本地制造步骤的恢复方式，减少重复调用模型带来的时间和成本。",
            "在 M4 阶段，我参与接入 Bambu X1C 的局域网连接验证、FTPS 上传和 MQTT 启动指令，并处理 AMS 多材料映射。双材料实测中，需要把逻辑颜色准确映射到真实料槽，并保证发送给打印机的 ams_mapping 是正确数组而不是嵌套字符串。最终流程能够在不手动操作 Bambu Studio 图形界面的情况下完成灰色到黄色的自动换料。对于启动打印这一不可逆步骤，我采用“一次发送、不自动重放”的策略；访问码只在需要时读取，不写入状态文件、日志或浏览器。",
        ],
    ),
    (
        "（四）本地网页、语音输入与任务管理",
        2,
        [
            "为了降低使用门槛，我开发和完善了本地“造物台”网页。用户可以输入文字或录制语音，语音通过 faster-whisper 在本机转写，转写结果仍可编辑后再提交。网页能够展示阶段进度、错误信息和交付文件，并提供暂停、继续、停止、失败重试、任务置顶以及安全条件下的再次打印功能。服务只监听 127.0.0.1，临时音频会在转写后删除，打印机密钥也不会下发到前端。",
            "我还补充了实体打印状态摘要与监测逻辑，将打印中、暂停、完成、告警、剩余时间和冷却状态转换为用户能理解的提示。暂停并不等于任务结束，因此再次打印前必须确认原任务已经结束、设备无告警并进入安全空闲状态。前端控制只在软件阶段边界生效，打印指令一旦发送，网页不会把“停止按钮”包装成实体急停，这一边界在界面和后端中都进行了明确处理。",
        ],
    ),
    (
        "（五）测试、文档与交付验证",
        2,
        [
            "除功能开发外，我持续编写单元测试和回归测试，覆盖需求解析、Schema 校验、模型提交、网格修复、归一化、切片配置、AMS 映射、文件下载、语音转写和打印监测等环节。对于真实设备测试，我记录请求编号、文件哈希、状态事件和验收结果，区分离线验证、真实切片和实体打印，避免把某一环境下的通过结论扩大为生产承诺。同时我整理了快速开始、自动打印工作流、模块规范和故障恢复记录，使项目能够被复现和继续维护。",
        ],
    ),
    (
        "三、遇到的困难及解决办法",
        1,
        [],
    ),
    (
        "（一）自然语言灵活，但工程参数必须确定",
        2,
        [
            "最大的矛盾是用户表达往往模糊，而制造任务需要确定参数。例如“做一个结实的支架”并没有给出载荷、安装接口和尺寸。如果直接让大模型猜测，结果可能看似完整却不安全。我将大模型限制在意图理解和信息提取范围，把单位换算、枚举值、缺失项和状态转换交给本地程序；工程信息不足时必须澄清，只有状态为 ready 或 complete 的任务才能进入 M2。这样既利用了模型的语言能力，又保留了工程系统的确定性。",
        ],
    ),
    (
        "（二）AI 模型“看起来正确”不等于可打印",
        2,
        [
            "生成模型常出现非封闭网格、重复面、微小悬空、支撑断层或尺度不一致。早期处理中过度依赖单一几何指标，容易出现误报或修复后外形变化。我改为分阶段检查：先确认文件和坐标，再检查网格拓扑及尺寸，必要时生成多个修复候选，并用体积、外廓和连通性限制修复幅度。切片前后均保留可回读的证据。对于 Windows 长路径导致外部工具失败的问题，我使用同目录短临时文件和原子替换，失败时保护原文件，从而提升了流程稳定性。",
        ],
    ),
    (
        "（三）外部工具与真实设备带来大量边界情况",
        2,
        [
            "项目横跨 Windows、WSL、GPU 模型环境、Bambu Studio、FTPS 和 MQTT，任一环节的路径、编码、超时或返回格式差异都会使流程中断。我通过统一路径解析、结构化错误码、阶段事件和可恢复状态，把“软件调用失败”与“几何不合格”分开处理，避免无意义地重新生成模型。在打印机侧，我坚持先做只读连接和预检，再上传，再读取新鲜的空闲状态，最后才发送一次启动指令；如果返回结果不明确，系统保留现场并要求人工确认。",
        ],
    ),
    (
        "（四）智能监测效果与产品承诺之间需要边界",
        2,
        [
            "我曾对 PrintGuard 和 YOLO 分类模型进行离线筛选，替代候选在项目留出集上的平衡准确率达到 73.33%，但这并不能代表 X1C 实时摄像头场景中的生产准确率。同时，X1C 本身已有原生 AI 监测能力，因此当前方案把自定义模型定位为辅助研究检测器，自动暂停和停止仍保持关闭。这个过程让我体会到，算法指标只是验证链的一部分，只有经过目标设备、真实材料和异常场景测试后，才能转化为可靠的产品能力。",
        ],
    ),
    (
        "四、收获与心得",
        1,
        [
            "通过本次实习，我最大的收获是建立了从 AI 原型到工程系统的完整认识。过去我更关注单个模型或功能是否能够运行，而在这个项目中，我开始重视模块契约、输入输出验证、状态机、幂等性、密钥保护和失败恢复。特别是在实体制造场景中，程序的“成功返回”并不等于产品一定可用，必须区分模型生成、切片成功、文件可交付、设备已启动和实物质量合格等不同层次。",
            "我也提升了跨领域学习能力。为了完成项目，我需要同时理解自然语言处理、JSON Schema、三维网格、坐标与单位、切片参数、支撑结构、网络协议、前后端交互和测试方法。遇到问题时，我逐渐形成了先保留证据、再缩小范围、最后修改并回归验证的习惯。真实打印和多材料换料让我认识到，工程开发不能只停留在屏幕上的结果，只有让软件、设备和材料共同完成闭环，才能发现真正有价值的问题。",
            "此外，我对“智能化”的理解也发生了变化。智能化不是把所有决策都交给 AI，而是把适合概率模型处理的理解与生成任务，和适合确定性程序处理的校验、约束及安全控制结合起来。一个可靠的系统不仅要能完成理想路径，还要在信息不足、外部工具失败、网络中断或实体结果未知时做出保守且可解释的选择。",
        ],
    ),
    (
        "五、不足与后续计划",
        1,
        [
            "目前系统仍存在局限：生成模型质量受模型能力和算力影响，工程零件的知识覆盖还不够完整；自动切片成功不能替代材料、喷嘴、平台和首层条件的人工确认；视觉监测尚未在 X1C 实时摄像头域完成充分验证。后续我计划扩充制造知识库和需求样本，完善不同材料与设备的配置管理，增加真实打印数据采集和故障案例回放，并研究 X1C 原生 AI 信号与自定义检测器的安全融合。在任何自动干预上线前，都应先完成离线评估、影子运行和人工审核。",
        ],
    ),
    (
        "六、总结",
        1,
        [
            "本次实习让我完成了一个从自然语言到 3D 打印的端到端系统实践。我不仅实现了需求解析、AI 建模、网格处理、自动切片、网页交互和设备接入，也在不断调试中理解了制造软件对准确性、可追溯性和安全边界的要求。未来我将继续完善真实设备验证和检测优化，使“所想即所得”从概念进一步走向稳定、可信、可复现的智能增材制造应用。",
        ],
    ),
]


def build() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_abstract(doc)

    for heading, level, paragraphs in SECTIONS:
        add_heading(doc, heading, level)
        for paragraph in paragraphs:
            add_body(doc, paragraph)

    props = doc.core_properties
    props.title = "自然语言驱动的智能化增材制造系统开发实习报告"
    props.subject = "物质创制——AI 融合 3D 打印项目"
    props.author = ""
    props.keywords = "自然语言处理, 生成式AI, 3D打印, 自动切片, Bambu X1C"
    props.comments = "根据项目仓库和产品材料整理的实习报告"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)

    content = "".join(
        [
            "本次实习围绕物质创制项目展开",
            *(heading for heading, _, _ in SECTIONS),
            *(paragraph for _, _, paragraphs in SECTIONS for paragraph in paragraphs),
        ]
    )
    chinese_chars = sum("\u4e00" <= char <= "\u9fff" for char in content)
    visible_chars = sum(not char.isspace() for char in content)
    print(f"OUTPUT={OUTPUT}")
    print(f"CHINESE_CHARACTERS={chinese_chars}")
    print(f"VISIBLE_CHARACTERS={visible_chars}")


if __name__ == "__main__":
    build()
